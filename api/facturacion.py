"""Facturación: filas por maniobra, carga del XML, ingresos en Gastos, y la
Auditoría del Excel que las compara.

Fase 1 de docs/planes/PLAN_MODULO_FINANZAS.md (rama main). El número entre
paréntesis, p. ej. (P11), remite a la pregunta de PREGUNTAS_MODULO_FINANZAS.md
que lo decidió.

Desde el 2026-10-15 hay dos tablas (usuario):
- `Factura` = FACTURACIÓN. Nace al asignar el folio, la liga el No. Factura y
  la rellena el XML. Es la que manda: Cobrar, Dashboards e Ingresos.
- `FacturaExcel` = AUDITORÍA. El Excel de siempre, solo para comparar.
"""
import csv
import io
import re
import unicodedata
import xml.etree.ElementTree as ET
from difflib import SequenceMatcher
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo

from django.core.exceptions import ObjectDoesNotExist
from django.db import transaction
from django.db.models import F, Max, Q, Sum
from django.db.models.functions import Coalesce
from django.db.models.functions import TruncWeek
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import mixins, serializers, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.pagination import PageNumberPagination
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from .db_context import get_db_alias
from .models import (Abono, Banco, ClientePrincipal, Empleado, Factura, FacturaExcel, Gasto, Maniobra, Pago,
                     PerfilUsuario)

# ── Quién entra ──────────────────────────────────────────────────────────────
# Cerrado por defecto (P45, P64): staff o uno de estos cargos EXACTOS. Sin
# perfil, sin empleado, empleado borrado u otro cargo → fuera. Al revés que el
# plan de roles por cargo, que deja pasar: aquí se trata de dinero.
CARGOS_FINANZAS = {'COMERCIAL', 'DIRECTOR GENERAL', 'DIRECTOR OPERATIVO',
                      'DIRECTORA COMERCIAL'}


# El rendimiento que calcula el sistema en Reportes de Viaje: staff y dirección,
# sin Comercial (usuario, 2026-10-13).
CARGOS_DIRECCION = {'DIRECTOR GENERAL', 'DIRECTOR OPERATIVO', 'DIRECTORA COMERCIAL'}


def cargo_de(usuario):
    """El cargo del empleado ligado al usuario, normalizado; '' sin empleado."""
    try:
        empleado = usuario.perfil.empleado
    except (PerfilUsuario.DoesNotExist, Empleado.DoesNotExist):
        return ''
    # strip/upper: empleados.cargo es texto libre editable a mano en pgAdmin.
    return (empleado.cargo or '').strip().upper() if empleado else ''


def staff_o_cargo(usuario, cargos):
    return usuario.is_staff or cargo_de(usuario) in cargos


def admin_de_sistemas(usuario):
    """Folios: staff Y cargo SISTEMAS, las dos (usuario, 2026-10-14). Al reves
    que el resto, ser admin no basta."""
    return usuario.is_staff and cargo_de(usuario) == 'SISTEMAS'


def puede_finanzas(usuario):
    return staff_o_cargo(usuario, CARGOS_FINANZAS)


def ve_rendimiento(usuario):
    return staff_o_cargo(usuario, CARGOS_DIRECCION)


# ── Serie + folio ────────────────────────────────────────────────────────────
EMPRESA_DE_SERIE = {'SEF': 'soluciones', 'S': 'fraba'}  # P29, P61
# Facturación empieza con las maniobras de 2026 (usuario, 2026-10-16; del
# 2026-10-15 al 16 fue julio): las anteriores no entran, para no tener que subir
# el XML de cada factura vieja.
INICIO_FACTURACION = date(2026, 1, 1)
# Las filas SIN XML de maniobras anteriores al segundo semestre no salen en el
# listado (usuario, 2026-10-30): no se les va a subir el XML. Se ocultan, no se
# borran, y si llega su XML la fila se rellena y vuelve a salir.
INICIO_XML_PENDIENTE = date(2026, 7, 1)

_PIEZA = re.compile(r'^([A-Z]*)(\d+)$')


def _folio(valor):
    """'0456', 456, 456.0 → '456'. None si no es un número entero."""
    try:
        numero = Decimal(str(valor).strip())
    except InvalidOperation:
        return None
    if numero != numero.to_integral_value() or numero < 0:
        return None
    return str(int(numero))


def facturas_de_no_factura(texto):
    """Las (serie, folio) que nombra una celda No. Factura de Maniobras.

    'I 154/ 153' y 'I 154 // 153' → [('I', '154'), ('I', '153')]: la diagonal,
    una o dos, separa facturas que se suman (P2), y el trozo sin serie hereda la
    del anterior. Se ignoran mayúsculas, espacios y guiones (P6). Lo que no es
    serie + folio ('EFECTIVO', un '201' suelto) no nombra nada (P5).
    """
    facturas, serie = [], ''
    for trozo in (texto or '').split('/'):
        pieza = re.sub(r'[\s-]', '', trozo).upper()
        m = _PIEZA.match(pieza)
        if not m:
            continue
        serie = m.group(1) or serie
        if serie:
            facturas.append((serie, _folio(m.group(2))))
    return facturas


# ── Ingresos en Gastos (P58) ─────────────────────────────────────────────────
def recalcular_ingresos(maniobra_ids, usuario):
    """Gasto.facturado = suma de las facturas ACTIVAS ligadas a la maniobra.

    Solo se llama con maniobras cuya liga acaba de cambiar. Una que se queda sin
    facturas se queda en blanco: el importe que tenía era de las facturas que
    ya no son suyas. Una maniobra cuya liga no cambia no se toca, así que un
    ingreso escrito a mano antes de Facturación sigue ahí.
    """
    for maniobra_id in set(filter(None, maniobra_ids)):
        total = (Factura.objects.filter(maniobra_id=maniobra_id, estado='activa')
                 .aggregate(t=Sum('total'))['t'])
        # update() y no save(): save() recalcula gastos_totales, que no cambia.
        # updated_at a mano porque update() no dispara auto_now, y el refresco
        # automático de la tabla de Gastos mira esa columna.
        Gasto.objects.filter(maniobra_id=maniobra_id).update(
            facturado=str(total) if total is not None else None,
            updated_by=usuario, updated_at=timezone.now(),
        )


def _vacia(factura):
    """Sin XML ni cobro capturado: solo el hueco de una factura por emitir. Es
    la única fila que se borra, y la misma condición la exige la política de
    borrado de la base (migración 0081)."""
    return factura.uuid is None and factura.fecha_cobro is None and not factura.clave_rastreo


def sincronizar_facturas(maniobra, usuario):
    """Deja las filas de Facturación de la maniobra como dice su No. Factura.

    Una fila por factura nombrada; sin ninguna nombrada, una vacía, que es la
    que nace al asignar el folio (usuario, 2026-10-15). Sin folio y sin filas
    aún no hay nada que facturar.

    Cada factura nombrada: si la maniobra ya la tiene, se queda; si su XML llegó
    antes y está sin ligar, se liga; si la tiene otra maniobra no se le quita
    (con un número mal tecleado, robársela a la correcta sería un error
    silencioso, P3); si no, se reaprovecha una fila vacía o se crea.
    Lo que sobra: la fila con XML o con cobro se suelta y queda pendiente de
    ligar; la vacía se borra, salvo que sea la única que le queda.
    Si todas las que nombra son de otra maniobra (una factura que cobró dos
    servicios, usuario 2026-10-23), no le queda ninguna: una vacía saldría en
    blanco para siempre, porque su XML ya rellenó la de la otra.
    """
    if not maniobra.fecha_pis or maniobra.fecha_pis < INICIO_FACTURACION:
        return  # anterior a Facturación, o sin fecha para saberlo: no se toca
    propias = list(Factura.objects.filter(maniobra=maniobra).order_by('id'))
    tiene_folio = (maniobra.folio or '').strip() or (maniobra.folio_2 or '').strip()
    if not propias and not tiene_folio:
        return
    nombradas = list(dict.fromkeys(facturas_de_no_factura(maniobra.no_factura)))
    tiene = {(f.serie, f.folio) for f in propias}
    sobran = [f for f in propias if (f.serie, f.folio) not in nombradas]
    dinero = False  # solo si se liga o suelta una con importe se tocan los Ingresos

    for serie, folio in nombradas:
        if (serie, folio) in tiene:
            continue
        mismas = Factura.objects.filter(serie=serie, folio=folio)
        suelta = mismas.filter(maniobra__isnull=True).first()
        if suelta:
            suelta.maniobra, suelta.updated_by = maniobra, usuario
            suelta.save(update_fields=['maniobra', 'updated_by', 'updated_at'])
            dinero |= suelta.total is not None
            continue
        if mismas.exists():
            continue  # es de otra maniobra
        datos = {'serie': serie, 'folio': folio, 'empresa': EMPRESA_DE_SERIE.get(serie, '')}
        hueco = next((f for f in sobran if _vacia(f)), None)
        if hueco:
            sobran.remove(hueco)
            for campo, valor in datos.items():
                setattr(hueco, campo, valor)
            hueco.updated_by = usuario
            hueco.save(update_fields=[*datos, 'updated_by', 'updated_at'])
        else:
            Factura.objects.create(maniobra=maniobra, created_by=usuario,
                                   updated_by=usuario, **datos)

    for f in sorted(sobran, key=_vacia):  # primero las que se sueltan
        if not _vacia(f):
            f.maniobra, f.updated_by = None, usuario
            f.save(update_fields=['maniobra', 'updated_by', 'updated_at'])
            dinero |= f.total is not None
        elif nombradas or Factura.objects.filter(maniobra=maniobra).exclude(pk=f.pk).exists():
            f.delete()
        elif f.serie or f.folio:
            f.serie = f.folio = f.empresa = ''
            f.updated_by = usuario
            f.save(update_fields=['serie', 'folio', 'empresa', 'updated_by', 'updated_at'])

    if tiene_folio and not nombradas and not Factura.objects.filter(maniobra=maniobra).exists():
        # Con folio y sin factura nombrada (recién puesto, o se soltaron todas):
        # la fila por emitir.
        Factura.objects.create(maniobra=maniobra, created_by=usuario, updated_by=usuario)
    if dinero:
        recalcular_ingresos([maniobra.pk], usuario)


# ── Lectura del archivo ──────────────────────────────────────────────────────
def _cabecera(texto):
    """' Fecha Emisión' → 'fecha emision'. Por nombre y no por posición: el
    ejemplo trae ' Serie' y ' Folio' con un espacio delante."""
    sin_acentos = unicodedata.normalize('NFKD', str(texto or ''))
    return ''.join(c for c in sin_acentos if not unicodedata.combining(c)).strip().lower()


COLUMNAS = {
    'tipo': 'tipo comprobante', 'nombre': 'nombre', 'rfc': 'rfc',
    'serie': 'serie', 'folio': 'folio', 'uuid': 'uuid',
    'fecha': 'fecha emision', 'subtotal': 'subtotal', 'iva': 'iva tras',
    'total': 'total', 'moneda': 'moneda',
}


def _filas(archivo):
    """Las filas del archivo como listas, cabecera incluida. xlsx, xls o csv (P12)."""
    nombre = (archivo.name or '').lower()
    if nombre.endswith('.xlsx'):
        hoja = load_workbook(archivo, read_only=True, data_only=True).worksheets[0]
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    if nombre.endswith('.xls'):
        import xlrd  # solo aquí: es la única vía que lo necesita
        libro = xlrd.open_workbook(file_contents=archivo.read())
        hoja = libro.sheet_by_index(0)
        return [[xlrd.xldate_as_datetime(c.value, libro.datemode)
                 if c.ctype == xlrd.XL_CELL_DATE else c.value
                 for c in hoja.row(i)] for i in range(hoja.nrows)]
    if nombre.endswith('.csv'):
        crudo = archivo.read()
        try:
            texto = crudo.decode('utf-8-sig')
        except UnicodeDecodeError:
            texto = crudo.decode('latin-1')  # el export de Excel en Windows
        return list(csv.reader(io.StringIO(texto)))
    raise ValueError('Formato no admitido: sube un .xlsx, .xls o .csv.')


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or '').strip()
    for formato, largo in (('%Y-%m-%d', 10), ('%d/%m/%Y', 10)):
        try:
            return datetime.strptime(texto[:largo], formato).date()
        except ValueError:
            pass
    return None


def _importe(valor, obligatorio=True):
    if valor in (None, ''):
        return None if obligatorio else Decimal('0')
    try:
        return Decimal(str(valor).replace(',', '').strip()).quantize(Decimal('0.01'))
    except InvalidOperation:
        return None


def cargar_facturas(archivo, usuario):
    """Crea las facturas del archivo. Devuelve el resumen para la pantalla.

    La carga no se detiene por una fila mala (P11): se salta y se apunta en
    `omitidas`. Todo el archivo va en una transacción: o entra lo válido
    entero, o nada si revienta algo inesperado.
    """
    filas = _filas(archivo)
    if not filas:
        raise ValueError('El archivo está vacío.')
    indice = {_cabecera(c): i for i, c in enumerate(filas[0])}
    faltan = [c for c in COLUMNAS.values() if c not in indice]
    if faltan:
        raise ValueError('Faltan columnas: ' + ', '.join(faltan) + '.')

    existentes = set(FacturaExcel.objects.values_list('uuid', flat=True))
    nuevas, omitidas = [], []
    for n, fila in enumerate(filas[1:], start=2):  # n = fila de Excel
        dato = {k: (fila[indice[c]] if indice[c] < len(fila) else None)
                for k, c in COLUMNAS.items()}
        if all(v in (None, '') for v in dato.values()):
            continue

        def omitir(motivo):
            omitidas.append({'fila': n, 'serie': str(dato['serie'] or '').strip(),
                             'folio': str(dato['folio'] or '').strip(), 'motivo': motivo})

        serie = str(dato['serie'] or '').strip().upper()
        uuid = str(dato['uuid'] or '').strip().lower()
        tipo = str(dato['tipo'] or '').strip().upper()
        moneda = str(dato['moneda'] or '').strip().upper()
        folio, fecha = _folio(dato['folio']), _fecha(dato['fecha'])
        subtotal, total = _importe(dato['subtotal']), _importe(dato['total'])
        iva = _importe(dato['iva'], obligatorio=False)

        if tipo != 'I':  # P66: solo comprobantes de ingreso
            omitir(f'Tipo de comprobante {tipo or "vacío"}, no es de ingreso')
        elif moneda != 'MXN':  # P17
            omitir(f'Moneda {moneda or "vacía"}, solo se admite MXN')
        elif serie not in EMPRESA_DE_SERIE:
            omitir(f'Serie desconocida: {serie or "vacía"}')
        elif not uuid:
            omitir('Sin UUID')
        elif uuid in existentes:  # P11: en la base o antes en el mismo archivo
            omitir('UUID repetido')
        elif not folio or not fecha or subtotal is None or total is None or iva is None:
            omitir('Folio, fecha o importes ilegibles')
        else:
            existentes.add(uuid)
            nuevas.append(FacturaExcel(
                empresa=EMPRESA_DE_SERIE[serie], serie=serie, folio=folio, uuid=uuid,
                nombre=str(dato['nombre'] or '').strip()[:255],
                rfc=str(dato['rfc'] or '').strip()[:20],
                fecha_emision=fecha, subtotal=subtotal, iva=iva, total=total,
                created_by=usuario, updated_by=usuario,
            ))

    # Solo Auditoría: ya no liga ni toca Ingresos, eso es de Facturación.
    FacturaExcel.objects.bulk_create(nuevas)
    return {'creadas': len(nuevas), 'omitidas': omitidas}


# ── Lectura del XML (CFDI) ───────────────────────────────────────────────────
MAX_XML = 2 * 1024 * 1024  # un CFDI con Carta Porte pesa decenas de KB


def _local(etiqueta):
    """'{http://www.sat.gob.mx/cfd/4}Receptor' → 'Receptor'. Por nombre local:
    así lee igual el CFDI 4.0 que el 3.3."""
    return etiqueta.rsplit('}', 1)[-1]


def _hijos(nodo, nombre):
    return [h for h in (nodo if nodo is not None else []) if _local(h.tag) == nombre]


def _primero(nodo, nombre):
    return next(iter(_hijos(nodo, nombre)), None)


def _unir(valores):
    return ' / '.join(dict.fromkeys(v.strip() for v in valores if v and v.strip()))


def leer_cfdi(contenido):
    """Los datos de un CFDI. ValueError con el motivo si no se puede leer.

    El SAT exige UTF-8 y un CFDI nunca trae DOCTYPE: el que lo trae se rechaza
    antes de parsear. Eso cierra las bombas de entidades y las entidades
    externas sin añadir defusedxml.
    """
    if len(contenido) > MAX_XML:
        raise ValueError('El archivo pasa de 2 MB')
    try:
        texto = contenido.decode('utf-8-sig').upper()
    except UnicodeDecodeError:
        raise ValueError('No está en UTF-8') from None
    if '<!DOCTYPE' in texto or '<!ENTITY' in texto:
        raise ValueError('No es un CFDI válido')
    try:
        raiz = ET.fromstring(contenido)
    except ET.ParseError:
        raise ValueError('No es un XML válido') from None
    if _local(raiz.tag) != 'Comprobante':
        raise ValueError('No es un CFDI')

    a = raiz.attrib
    receptor = _primero(raiz, 'Receptor')
    impuestos = _primero(raiz, 'Impuestos')  # el del comprobante, no el de cada concepto
    conceptos = _hijos(_primero(raiz, 'Conceptos'), 'Concepto')
    timbre = next((n for n in raiz.iter() if _local(n.tag) == 'TimbreFiscalDigital'), None)
    de = lambda nodo, clave: (nodo.get(clave) or '') if nodo is not None else ''
    return {
        'tipo': a.get('TipoDeComprobante', '').strip().upper(),
        'moneda': a.get('Moneda', '').strip().upper(),
        'serie': a.get('Serie', '').strip().upper(),
        'folio': _folio(a.get('Folio', '')),
        'fecha': _fecha(a.get('Fecha')),
        'subtotal': _importe(a.get('SubTotal')),
        'iva': _importe(de(impuestos, 'TotalImpuestosTrasladados'), obligatorio=False),
        # El Total del CFDI: con IVA y sin la retención (usuario, 2026-10-15).
        'total': _importe(a.get('Total')),
        'uuid': de(timbre, 'UUID').strip().lower(),
        'nombre': de(receptor, 'Nombre').strip()[:255],
        'rfc': de(receptor, 'Rfc').strip().upper()[:20],
        'no_identificacion': _unir(c.get('NoIdentificacion') for c in conceptos),
        'descripcion': _unir(c.get('Descripcion') for c in conceptos),
    }


# Lo que escribe el XML. Sin `estado` ni la liga: el rol estándar no tiene
# permiso sobre esas columnas (migración 0081).
CAMPOS_XML = ['empresa', 'serie', 'folio', 'uuid', 'nombre', 'rfc', 'fecha_emision',
              'subtotal', 'iva', 'total', 'moneda', 'no_identificacion', 'descripcion']


def cargar_xml(archivos, usuario):
    """Rellena con cada XML la fila de Facturación de su serie y folio.

    Sin fila que lo espere, crea una sin maniobra (pendiente de ligar) que se
    liga sola cuando alguien capture ese No. Factura (usuario, 2026-10-15).
    Como el Excel, un archivo malo no para la carga: se apunta en `omitidas`.
    """
    omitidas, completadas, sin_maniobra, maniobras = [], 0, 0, set()
    vistos = set(Factura.objects.exclude(uuid=None).values_list('uuid', flat=True))
    with transaction.atomic(using=get_db_alias()):
        for archivo in archivos:
            def omitir(motivo, d=None):
                omitidas.append({'archivo': archivo.name, 'serie': (d or {}).get('serie', ''),
                                 'folio': (d or {}).get('folio') or '', 'motivo': motivo})
            try:
                d = leer_cfdi(archivo.read(MAX_XML + 1))
            except ValueError as e:
                omitir(str(e))
                continue
            if d['tipo'] != 'I':  # P66: solo comprobantes de ingreso
                omitir(f'Tipo de comprobante {d["tipo"] or "vacío"}, no es de ingreso', d)
                continue
            if d['moneda'] != 'MXN':  # P17
                omitir(f'Moneda {d["moneda"] or "vacía"}, solo se admite MXN', d)
                continue
            if d['serie'] not in EMPRESA_DE_SERIE:
                omitir(f'Serie desconocida: {d["serie"] or "vacía"}', d)
                continue
            if not d['uuid']:
                omitir('Sin UUID: no está timbrado', d)
                continue
            if d['uuid'] in vistos:
                # Las cargadas con el Excel antes de la 0081 tienen UUID pero no
                # estos dos campos, que solo trae el XML: se completan y nada más.
                if Factura.objects.filter(uuid=d['uuid'], no_identificacion='', descripcion='').update(
                        no_identificacion=d['no_identificacion'], descripcion=d['descripcion'],
                        updated_by=usuario, updated_at=timezone.now()):
                    completadas += 1
                else:
                    omitir('Ese XML ya estaba cargado', d)
                continue
            if not d['folio'] or not d['fecha'] or d['subtotal'] is None or d['total'] is None:
                omitir('Folio, fecha o importes ilegibles', d)
                continue
            mismas = Factura.objects.filter(serie=d['serie'], folio=d['folio'])
            huecos = list(mismas.filter(uuid__isnull=True)[:2])
            if len(huecos) > 1:
                omitir('Varias maniobras tienen ese No. Factura: corrígelo en Control de viajes', d)
                continue
            if not huecos and mismas.exists():
                omitir('Ya hay otro XML con esa serie y folio', d)
                continue

            vistos.add(d['uuid'])
            f = huecos[0] if huecos else Factura(created_by=usuario)
            f.empresa, f.fecha_emision, f.moneda = EMPRESA_DE_SERIE[d['serie']], d['fecha'], 'MXN'
            for campo in CAMPOS_XML:
                if campo in d:
                    setattr(f, campo, d[campo])
            f.updated_by = usuario
            if f.pk:
                f.save(update_fields=CAMPOS_XML + ['updated_by', 'updated_at'])
                completadas += 1
                maniobras.add(f.maniobra_id)
            else:
                f.save()
                sin_maniobra += 1
        recalcular_ingresos(maniobras, usuario)
    return {'completadas': completadas, 'sin_maniobra': sin_maniobra, 'omitidas': omitidas}


# ── Auditoría: Excel contra Facturación ──────────────────────────────────────
def _igual(valor):
    """Para comparar sin que cuenten mayúsculas ni espacios de más."""
    return ' '.join(str(valor if valor is not None else '').split()).upper()


COMPARADOS = (('nombre', 'Cliente'), ('rfc', 'RFC'), ('fecha_emision', 'Fecha de emisión'),
              ('total', 'Total'), ('uuid', 'UUID'), ('estado', 'Estado'))


def _folio_de_maniobra(factura):
    try:
        return factura.maniobra.folio if factura.maniobra_id else None
    except ObjectDoesNotExist:
        return None


def comparar():
    """Cada factura del Excel contra su pareja de Facturación (misma serie y
    folio), y las de Facturación con XML que el Excel no trae.

    La maniobra y el vencimiento de una fila de Auditoría son los de su pareja:
    el Excel no se liga por su cuenta. El UUID se compara pero no sale.
    """
    pareja = {}
    for f in (Factura.objects.exclude(serie='')
              .select_related('maniobra__cliente_fk__cliente_principal').order_by('id')):
        clave = (f.serie, f.folio)
        if clave not in pareja or (pareja[clave].uuid is None and f.uuid):
            pareja[clave] = f  # si hay dos, la que ya tiene XML

    filas, en_excel = [], set()
    for e in FacturaExcel.objects.all():
        clave = (e.serie, e.folio)
        en_excel.add(clave)
        f = pareja.get(clave)
        if f is None:
            diferencias = ['Falta en Facturación']
        elif f.uuid is None:
            diferencias = ['Falta el XML en Facturación']
        else:
            diferencias = [etiqueta for campo, etiqueta in COMPARADOS
                           if _igual(getattr(e, campo)) != _igual(getattr(f, campo))]
        principal = principal_de(f) if f else None
        filas.append({
            'id': e.id, 'empresa': e.empresa, 'serie': e.serie, 'folio': e.folio,
            'nombre': e.nombre, 'rfc': e.rfc, 'fecha_emision': e.fecha_emision,
            'total': _dinero(e.total), 'estado': e.estado,
            'vencimiento': (e.fecha_emision + timedelta(days=principal.dias_credito)
                            if principal else None),
            'maniobra': f.maniobra_id if f else None,
            'maniobra_folio': _folio_de_maniobra(f) if f else None,
            'diferencias': diferencias,
        })
    solo = [{'id': f.id, 'serie': f.serie, 'folio': f.folio, 'nombre': f.nombre,
             'total': _dinero(f.total)}
            for clave, f in pareja.items() if f.uuid and clave not in en_excel]
    return {'facturas': filas, 'solo_en_facturacion': solo}


# ── Cuentas por cobrar (Fase 3) ──────────────────────────────────────────────
# Tramos de antigüedad por días VENCIDOS (P28). `por_vencer` son las que aún no
# llegan a su vencimiento.
# El "hoy" de la antigüedad es el de México, no el del servidor (TIME_ZONE =
# 'UTC'): con UTC, desde las 18:00 locales todas las facturas envejecerían un
# día antes de tiempo. Misma zona que _ZONA_OPERACION en views.py.
ZONA_OPERACION = ZoneInfo('America/Mexico_City')

TRAMOS = [('por_vencer', None), ('0_30', 30), ('31_60', 60), ('61_90', 90), ('mas_90', None)]


def _dinero(valor):
    """Como el resto de la API (DecimalField del serializer): texto con dos
    decimales. Un Decimal suelto en un Response saldría como float en el JSON."""
    return str(Decimal(valor or 0).quantize(Decimal('0.01')))


def principal_de(factura):
    """El cliente principal de la factura, vía maniobra → dirección. Si falta
    algún eslabón, el principal con el mismo nombre que el receptor del XML."""
    try:
        direccion = factura.maniobra.cliente_fk if factura.maniobra_id else None
        principal = direccion.cliente_principal if direccion else None
    except (AttributeError, ObjectDoesNotExist):
        # Maniobra o dirección borradas a mano (sin FK en la base): sin cliente.
        principal = None
    return principal or principal_por_nombre(factura.nombre)


def principal_por_nombre(nombre):
    """El principal cuyo nombre, sin tipo de sociedad, es IGUAL al del receptor
    (usuario, 2026-10-26): el CFDI 4.0 trae "LOGISTICS PLUS" y el catálogo
    "LOGISTICS PLUS S DE RL DE CV". Igual y no parecido: dos empresas de nombre
    casi idéntico no se mezclan. None si no hay uno solo.
    ponytail: una consulta por factura sin eslabón; basta para decenas de
    clientes, pasar el índice por parámetro si los dashboards se vuelven lentos."""
    buscado = _razon_social(nombre)
    if not buscado:
        return None
    iguales = [p for p in ClientePrincipal.objects.all() if _razon_social(p.nombre) == buscado]
    return iguales[0] if len(iguales) == 1 else None


def vencimiento(factura):
    """Emisión + días de crédito del principal, en días naturales (P28, P36).
    Calculado y no guardado: cambiar los días de un cliente mueve sus facturas.
    None sin maniobra ligada o sin cliente principal: no hay días que aplicar."""
    principal = principal_de(factura) if factura.fecha_emision else None
    return factura.fecha_emision + timedelta(days=principal.dias_credito) if principal else None


def tramo(dias_vencida):
    if dias_vencida < 0:
        return 'por_vencer'
    for clave, tope in TRAMOS[1:-1]:
        if dias_vencida <= tope:
            return clave
    return 'mas_90'


def estatus(dias_vencida):
    """El de Estados de cuenta (usuario, 2026-10-21). Vencida desde el MISMO día
    que vence, como en su Excel; por vencer, la semana anterior. None sin
    vencimiento (sin cliente principal)."""
    if dias_vencida is None:
        return None
    return 'vencida' if dias_vencida >= 0 else 'por_vencer' if dias_vencida >= -7 else 'en_tiempo'


def cuentas_por_cobrar(hoy):
    """Facturas activas sin cobrar, con su antigüedad.

    Una factura sin maniobra ligada o cuya dirección no tiene cliente principal
    no tiene días de crédito: se lista aparte y queda FUERA de la antigüedad
    hasta que se ligue (decidido con el usuario el 2026-09-30).
    """
    antiguedad = {clave: Decimal('0') for clave, _ in TRAMOS}
    fuera = {'pendiente_de_ligar': Decimal('0'), 'sin_cliente_principal': Decimal('0')}
    filas = []
    qs = (Factura.objects.filter(estado='activa', fecha_cobro__isnull=True, total__isnull=False)
          .select_related('maniobra__cliente_fk__cliente_principal')
          .order_by('fecha_emision', 'id'))
    for f in qs:
        principal, venc = principal_de(f), vencimiento(f)
        # Lo que falta por cobrar, no el total (usuario, 2026-10-22).
        saldo = f.total - f.abonado
        if venc is None:
            situacion = 'pendiente_de_ligar' if not f.maniobra_id else 'sin_cliente_principal'
            fuera[situacion] += saldo
            dias, clave = None, None
        else:
            situacion = 'con_credito'
            dias = (hoy - venc).days
            clave = tramo(dias)
            antiguedad[clave] += saldo
        filas.append({
            'id': f.id, 'empresa': f.empresa, 'serie': f.serie, 'folio': f.folio,
            'nombre': f.nombre, 'fecha_emision': f.fecha_emision, 'saldo': _dinero(saldo),
            'cliente_principal': principal.nombre if principal else None,
            'cliente_principal_id': principal.id if principal else None,
            'no_identificacion': f.no_identificacion, 'descripcion': f.descripcion,
            'vencimiento': venc, 'dias_vencida': dias, 'tramo': clave,
            'estatus': estatus(dias), 'situacion': situacion,
        })
    return {'facturas': filas,
            'antiguedad': {k: _dinero(v) for k, v in antiguedad.items()},
            'fuera_de_antiguedad': {k: _dinero(v) for k, v in fuera.items()},
            'total': _dinero(sum(antiguedad.values()) + sum(fuera.values()))}


def cobranza_semanal(desde=None, hasta=None):
    """Lo COBRADO por semana (lunes) de la fecha del pago: el monto de los
    pagos activos de Compensación (usuario, 2026-10-22). Antes era lo
    facturado por semana de emisión (P67, P70)."""
    qs = Pago.objects.filter(estado='activa')
    if desde:
        qs = qs.filter(fecha__gte=desde)
    if hasta:
        qs = qs.filter(fecha__lte=hasta)
    return [{'semana': fila['semana'], 'total': _dinero(fila['total'])}
            for fila in (qs.annotate(semana=TruncWeek('fecha'))
                         .values('semana').annotate(total=Sum('monto')).order_by('semana'))]


# ── API ──────────────────────────────────────────────────────────────────────
class FacturaSerializer(serializers.ModelSerializer):
    # El folio de la maniobra ligada (P15). El UUID no sale nunca (P16).
    maniobra_folio = serializers.CharField(source='maniobra.folio', default=None,
                                           read_only=True)
    # La fecha de la factura sin XML, para el orden (usuario, 2026-10-28).
    fecha_pis = serializers.DateField(source='maniobra.fecha_pis', default=None, read_only=True)
    vencimiento = serializers.SerializerMethodField()
    saldo_pendiente = serializers.SerializerMethodField()

    def get_vencimiento(self, factura):
        return vencimiento(factura)

    def get_saldo_pendiente(self, factura):
        return None if factura.total is None else _dinero(factura.total - factura.abonado)

    class Meta:
        model = Factura
        fields = ['id', 'empresa', 'serie', 'folio', 'nombre', 'rfc',
                  'fecha_emision', 'subtotal', 'iva', 'total', 'estado',
                  'no_identificacion', 'descripcion', 'fecha_cobro', 'clave_rastreo',
                  'maniobra', 'maniobra_folio', 'fecha_pis', 'vencimiento', 'saldo_pendiente']
        read_only_fields = fields


class CancelableMixin:
    """Cancelar y reactivar: solo staff (P65), que escribe como superusuario."""

    def _cambiar_estado(self, request, estado):
        if not request.user.is_staff:
            raise PermissionDenied('Solo un administrador puede cancelar o reactivar facturas.')
        factura = self.get_object()
        with transaction.atomic(using=get_db_alias()):
            factura.estado = estado
            factura.updated_by = request.user.username
            factura.save(update_fields=['estado', 'updated_by', 'updated_at'])
            if isinstance(factura, Factura):
                recalcular_ingresos([factura.maniobra_id], request.user.username)
        return Response(self.get_serializer(factura).data)

    @action(detail=True, methods=['post'])
    def cancelar(self, request, pk=None):
        return self._cambiar_estado(request, 'cancelada')

    # ponytail: reactivar no se pidió; existe para deshacer una cancelación por error.
    @action(detail=True, methods=['post'])
    def reactivar(self, request, pk=None):
        return self._cambiar_estado(request, 'activa')


# ── Las dos tablas de la pantalla, de 100 en 100 (usuario, 2026-10-16) ──────
# La búsqueda y las fechas se aplican AQUÍ y sobre todas las filas: con la
# página de 100 en el navegador, filtrar allí solo vería lo ya cargado. Es la
# misma regla que utils/buscar.mjs y utils/filtroFechas.mjs del frontend (sin
# acentos ni mayúsculas, -palabra excluye, "frase" exacta; fecha, rango en
# cualquier orden y mes), sobre los valores de la fila tal como salen en el JSON.
# ponytail: filtra en Python sobre la lista entera (cientos o pocos miles de
# facturas); si crece mucho más, pasarlo a SQL (unaccent) y paginar el queryset.
class PaginaDeCien(PageNumberPagination):
    page_size = 100


_TERMINO = re.compile(r'-?"[^"]*"|\S+')


def _texto_buscable(valor):
    if isinstance(valor, (list, tuple)):
        valor = ' '.join(str(v) for v in valor)
    return _cabecera('' if valor is None else valor)


def buscar(filas, consulta):
    incluir, excluir = [], []
    for bruto in _TERMINO.findall(consulta or ''):
        negado = bruto.startswith('-')
        termino = _cabecera((bruto[1:] if negado else bruto).replace('"', ''))
        if termino:  # un '-' suelto o unas comillas vacías no filtran
            (excluir if negado else incluir).append(termino)
    if not incluir and not excluir:
        return filas

    def casa(fila):
        valores = [_texto_buscable(v) for v in fila.values()]
        hay = lambda t: any(t in v for v in valores)
        return all(map(hay, incluir)) and not any(map(hay, excluir))
    return [f for f in filas if casa(f)]


def en_fechas(valor, p):
    """?fecha= ?desde=&hasta= (las dos, en cualquier orden) ?mes=AAAA-MM."""
    fecha, desde, hasta, mes = (p.get(k) for k in ('fecha', 'desde', 'hasta', 'mes'))
    if not fecha and not (desde and hasta) and not mes:
        return True
    dia = str(valor)[:10] if valor else ''
    if not dia or (fecha and dia != fecha):
        return False
    if desde and hasta and not (min(desde, hasta) <= dia <= max(desde, hasta)):
        return False
    return not mes or dia.startswith(mes)


def filtrar_tabla(filas, p):
    return [f for f in buscar(filas, p.get('search')) if en_fechas(f.get('fecha_emision'), p)]


def ordenar_facturas(filas, orden, una_empresa):
    """De la más nueva a la más vieja (usuario, 2026-10-28).
    - ?orden=fecha (por defecto): las que aún no tienen XML arriba, por la fecha
      PIS de su maniobra; debajo las cargadas, por fecha de emisión.
    - ?orden=factura con una empresa: por número de folio (S 2000, S 1999…).
      En Todas, S y SEF no comparten numeración: se intercalan por fecha
      (emisión, o PIS sin XML) y el folio desempata dentro del mismo día.
    """
    fecha = lambda f: str(f['fecha_emision'] or f['fecha_pis'] or '')[:10]
    folio = lambda f: int(f['folio']) if f['folio'].isdigit() else -1
    if orden == 'factura':
        clave = folio if una_empresa else (lambda f: (fecha(f), folio(f)))
    else:
        clave = lambda f: (f['fecha_emision'] is None, fecha(f))
    return sorted(filas, key=clave, reverse=True)


class FacturaViewSet(CancelableMixin, mixins.ListModelMixin, viewsets.GenericViewSet):
    """FACTURACIÓN. Solo lectura salvo el XML y el cobro: las filas las crea la
    maniobra y las rellena el XML; nadie teclea importes (P14, P65)."""
    serializer_class = FacturaSerializer
    pagination_class = PaginaDeCien

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        # `acceso` responde a cualquiera: es como el frontend sabe si pintar
        # la tarjeta. Todo lo demás, solo quien puede.
        if self.action != 'acceso' and not puede_finanzas(request.user):
            raise PermissionDenied('Facturación es solo para dirección, comercial y administradores.')

    def get_queryset(self):
        qs = Factura.objects.select_related('maniobra__cliente_fk__cliente_principal')
        estado = self.request.query_params.get('estado')
        return qs.filter(estado=estado) if estado in ('activa', 'cancelada') else qs

    def list(self, request, *args, **kwargs):
        qs = self.get_queryset().exclude(uuid__isnull=True, maniobra__fecha_pis__lt=INICIO_XML_PENDIENTE)
        filas = self.get_serializer(qs, many=True).data
        # Una factura que cobró dos servicios se liga a uno; los demás que la
        # nombran salen al lado para que no parezca que no se facturaron.
        nombran = {}
        for m_id, folio, no_factura in (Maniobra.objects.filter(fecha_pis__gte=INICIO_FACTURACION)
                                        .exclude(no_factura__isnull=True).exclude(no_factura='')
                                        .values_list('id', 'folio', 'no_factura')):
            for clave in dict.fromkeys(facturas_de_no_factura(no_factura)):
                nombran.setdefault(clave, []).append((m_id, folio or f'Maniobra {m_id}'))
        for f in filas:
            f['comparte_con'] = [folio for m_id, folio in nombran.get((f['serie'], f['folio']), [])
                                 if m_id != f['maniobra']]
        p = request.query_params
        filas = filtrar_tabla(filas, p)
        empresa = p.get('empresa') if p.get('empresa') in ('fraba', 'soluciones') else None
        if empresa:
            filas = [f for f in filas if f['empresa'] == empresa]
        pagina = self.paginate_queryset(ordenar_facturas(filas, p.get('orden'), empresa))
        return self.get_paginated_response(pagina)

    @action(detail=False, methods=['get'])
    def acceso(self, request):
        return Response({'ver': puede_finanzas(request.user),
                         'auditoria': request.user.is_staff,
                         'cancelar': request.user.is_staff})

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser])
    def xml(self, request):
        archivos = request.FILES.getlist('archivos')
        if not archivos:
            return Response({'detail': 'Falta el XML.'}, status=status.HTTP_400_BAD_REQUEST)
        return Response(cargar_xml(archivos, request.user.username), status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='cuentas-por-cobrar')
    def cuentas_por_cobrar(self, request):
        return Response(cuentas_por_cobrar(datetime.now(ZONA_OPERACION).date()))

    @action(detail=False, methods=['get'], url_path='cobranza-semanal')
    def cobranza_semanal(self, request):
        desde = _fecha(request.query_params.get('desde')) if request.query_params.get('desde') else None
        hasta = _fecha(request.query_params.get('hasta')) if request.query_params.get('hasta') else None
        return Response(cobranza_semanal(desde, hasta))

    @action(detail=False, methods=['get'])
    def clientes(self, request):
        """Los clientes de Compensación: los RFC de las facturas, con su nombre."""
        nombres = {}
        for f in Factura.objects.exclude(rfc='').order_by('-fecha_emision').values('rfc', 'nombre'):
            nombres.setdefault(f['rfc'], f['nombre'])  # el nombre más reciente
        return Response(sorted(({'rfc': r, 'nombre': n} for r, n in nombres.items()),
                               key=lambda c: c['nombre']))

    @action(detail=False, methods=['get'], url_path='buscar-compensacion')
    def buscar_compensacion(self, request):
        """El buscador de Compensación: TODAS las facturas (también canceladas y
        sin XML, que salen pero no se abonan), del cliente elegido si lo hay."""
        # ponytail: filtra en Python como la tabla de arriba; a SQL si crece.
        qs = Factura.objects.exclude(folio='').annotate(
            ultima=Max('abonos__parcialidad', filter=Q(abonos__estado='activa')))
        rfc = request.query_params.get('rfc')
        if rfc:
            qs = qs.filter(rfc=rfc)
        filas = [fila_compensacion(f, f.ultima) for f in qs]
        return Response(buscar(filas, request.query_params.get('search'))[:50])


def fila_compensacion(f, ultima=None):
    """Una factura tal como la usa Compensación. `ultima`: su última
    parcialidad activa, para sugerir la siguiente."""
    return {
        'id': f.id, 'serie': f.serie, 'folio': f.folio, 'fecha_emision': f.fecha_emision,
        'rfc': f.rfc, 'nombre': f.nombre, 'estado': f.estado,
        'total': None if f.total is None else _dinero(f.total),
        'abonado': _dinero(f.abonado),
        'saldo_pendiente': None if f.total is None else _dinero(f.total - f.abonado),
        'fecha_cobro': f.fecha_cobro,
        'abonable': f.estado == 'activa' and f.total is not None,
        'siguiente_parcialidad': (ultima or 0) + 1,
    }


def _fila_al_dia(factura):
    ultima = factura.abonos.filter(estado='activa').aggregate(m=Max('parcialidad'))['m']
    return fila_compensacion(factura, ultima)


# ── El CSV de movimientos del banco (usuario, 2026-10-22) ─────────────────────
# Formato de docs/planes/20261022_MovimientosCheque.csv (rama main). Cada celda
# de texto viene entre ' (marca de Excel; a veces solo delante) y rellena de
# espacios; la fecha es ddmmaaaa ('02102026' = 2 de octubre de 2026). Solo es pago lo que dice ABONO TRANSFERENCIA SPEI: el resto
# (intereses, cargos) se omite.
COLUMNAS_CSV = {'descripcion': 'descripcion', 'fecha': 'fecha', 'monto': 'importe',
                'clave_rastreo': 'clave de rastreo', 'banco': 'banco participante',
                'nombre': 'nombre ordenante', 'rfc': 'rfc ordenante'}
ES_PAGO = 'abono transferencia spei'


def _celda(valor):
    return (valor or '').strip().strip("'").strip()


def leer_movimientos(contenido):
    """Los pagos del CSV, sin guardar nada. Devuelve (pagos, omitidos, errores)."""
    try:
        texto = contenido.decode('utf-8-sig')
    except UnicodeDecodeError:
        texto = contenido.decode('latin-1')
    lector = csv.reader(io.StringIO(texto))
    cabecera = {_cabecera(c): i for i, c in enumerate(next(lector, []))}
    faltan = [c for c in COLUMNAS_CSV.values() if c not in cabecera]
    if faltan:
        raise ValueError(f"Al CSV le faltan columnas: {', '.join(faltan)}.")
    pagos, omitidos, errores = [], 0, []
    for n, fila in enumerate(lector, start=2):
        if not any(c.strip() for c in fila):
            continue
        v = {k: _celda(fila[cabecera[c]]) if cabecera[c] < len(fila) else ''
             for k, c in COLUMNAS_CSV.items()}
        if _cabecera(v.pop('descripcion')) != ES_PAGO:
            omitidos += 1
            continue
        try:
            v['fecha'] = datetime.strptime(v['fecha'], '%d%m%Y').date()
            v['monto'] = Decimal(v['monto'].replace(',', ''))
        except (ValueError, InvalidOperation):
            errores.append(f'Fila {n}: la fecha o el importe no se entienden.')
            continue
        if v['monto'] <= 0 or not re.fullmatch(r'[A-Za-z0-9]{1,50}', v['clave_rastreo']):
            errores.append(f'Fila {n}: sin importe o sin clave de rastreo válida.')
            continue
        v['clave_rastreo'] = v['clave_rastreo'].upper()
        v['rfc'] = v['rfc'].upper()[:20]
        v['banco'] = v['banco'][:80]
        pagos.append(v)
    return pagos, omitidos, errores


# El nombre del ordenante se compara de forma APROXIMADA con los clientes
# principales (usuario, 2026-10-22: "LOGISTICS PLUS" vs "LOGISTIC PLUS"). Sin el
# tipo de sociedad: "S DE RL DE CV" haría parecidas a dos empresas distintas.
_SOCIEDAD = re.compile(r'\b(S ?A ?P ?I|S ?A ?B|S ?DE ?R ?L|S ?EN ?C|S ?A|S ?C|A ?C|DE ?C ?V|DE)\b')
PARECIDO = 0.9


def _razon_social(nombre):
    limpio = re.sub(r'[^A-Z0-9 ]', ' ', _cabecera(nombre).upper())
    return ' '.join(_SOCIEDAD.sub(' ', limpio).split())


def cliente_parecido(nombre, principales):
    """El cliente principal cuyo nombre se parece al del ordenante, o None.
    ponytail: difflib sobre la lista entera; basta para decenas de clientes."""
    buscado = _razon_social(nombre)
    if not buscado:
        return None
    mejor = max(principales, default=None,
                key=lambda p: SequenceMatcher(None, buscado, _razon_social(p)).ratio())
    if mejor and SequenceMatcher(None, buscado, _razon_social(mejor)).ratio() >= PARECIDO:
        return mejor
    return None


def revisar_pagos(pagos):
    """Marca cada pago: repetido (su clave ya está en un pago activo o sale dos
    veces), cliente principal parecido y si su RFC tiene facturas."""
    claves = [p['clave_rastreo'] for p in pagos if p['clave_rastreo']]
    guardadas = set(Pago.objects.filter(estado='activa', clave_rastreo__in=claves)
                    .values_list('clave_rastreo', flat=True))
    con_facturas = set(Factura.objects.filter(rfc__in={p['rfc'] for p in pagos if p['rfc']})
                       .values_list('rfc', flat=True))
    principales = list(ClientePrincipal.objects.values_list('nombre', flat=True))
    vistas = set()
    for p in pagos:
        clave = p['clave_rastreo']
        p['repetido'] = bool(clave) and (clave in guardadas or clave in vistas)
        vistas.add(clave)
        p['cliente_principal'] = cliente_parecido(p['nombre'], principales)
        p['rfc_con_facturas'] = p['rfc'] in con_facturas
    return pagos


# ── Compensación (usuario, 2026-10-22) ───────────────────────────────────────
# Un PAGO es lo que llegó al banco (su monto es la cobranza); se reparte a mano
# en ABONOS, uno por factura. Abonar o pagar de más se acepta: la pantalla pide
# confirmación antes.
def mover_abonado(factura, delta, pago, usuario):
    """Suma (o resta, al cancelar) al abonado. Cubrir el total marca la factura
    cobrada con la fecha y la clave del pago que lo cubre; bajar de él la
    desmarca. Es la ÚNICA vía al cobro: ya no se captura a mano."""
    factura.abonado += delta
    if factura.abonado < factura.total:
        factura.fecha_cobro, factura.clave_rastreo = None, ''
    elif factura.fecha_cobro is None:
        factura.fecha_cobro, factura.clave_rastreo = pago.fecha, pago.clave_rastreo
    factura.updated_by = usuario
    factura.save(update_fields=['abonado', 'fecha_cobro', 'clave_rastreo',
                                'updated_by', 'updated_at'])


class PagoSerializer(serializers.ModelSerializer):
    clave_rastreo = serializers.RegexField(
        r'^[A-Za-z0-9]*$', max_length=50, required=False, allow_blank=True,
        error_messages={'invalid': 'La clave de rastreo solo admite letras y números.'})
    monto = serializers.DecimalField(max_digits=14, decimal_places=2, min_value=Decimal('0.01'))
    aplicado = serializers.SerializerMethodField()

    class Meta:
        model = Pago
        fields = ['id', 'fecha', 'monto', 'clave_rastreo', 'banco', 'rfc', 'nombre',
                  'estado', 'aplicado', 'created_by', 'created_at']
        read_only_fields = ['estado', 'created_by', 'created_at']

    def get_aplicado(self, pago):
        aplicado = getattr(pago, 'aplicado', None)
        if aplicado is None:
            aplicado = pago.abonos.filter(estado='activa').aggregate(t=Sum('monto'))['t']
        return _dinero(aplicado)

    def validate_clave_rastreo(self, clave):
        return (clave or '').upper()


class HistorialPagoSerializer(PagoSerializer):
    """HISTORIAL DE PAGOS: el pago con las facturas a las que se abonó."""
    abonos = serializers.SerializerMethodField()

    class Meta(PagoSerializer.Meta):
        fields = PagoSerializer.Meta.fields + ['abonos']

    def get_abonos(self, pago):
        return [{'id': a.id, 'factura': f'{a.factura.serie} {a.factura.folio}'.strip(),
                 'nombre': a.factura.nombre, 'monto': _dinero(a.monto),
                 'parcialidad': a.parcialidad, 'estado': a.estado}
                for a in pago.abonos.all()]


class _SoloFinanzas:
    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not puede_finanzas(request.user):
            raise PermissionDenied('Compensación es solo para dirección, comercial y administradores.')


class PagoViewSet(_SoloFinanzas, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                  viewsets.GenericViewSet):
    """COMPENSACIÓN, los pagos. Se leen del CSV o se capturan a mano en una
    lista y se guardan juntos; cancelar es solo staff (P65) y cancela también
    sus abonos. Sin borrado."""
    serializer_class = PagoSerializer
    queryset = Pago.objects.all()

    def list(self, request, *args, **kwargs):
        """El HISTORIAL: todos los pagos guardados, también los cancelados, del
        más reciente al más viejo. Lo aplicado se calcula al pedirlo.
        ponytail: sin paginar; cientos de pagos al año. Paginar si crece."""
        qs = (Pago.objects.annotate(aplicado=Coalesce(
                  Sum('abonos__monto', filter=Q(abonos__estado='activa')), Decimal('0')))
              .prefetch_related('abonos__factura').order_by('-fecha', '-id'))
        return Response(HistorialPagoSerializer(qs, many=True).data)

    @action(detail=False, methods=['post'], url_path='leer-csv',
            parser_classes=[MultiPartParser, FormParser])
    def leer_csv(self, request):
        """Solo LEE: lo que entraría, marcado. Nada se guarda aquí."""
        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'detail': 'Falta el CSV.'}, status=status.HTTP_400_BAD_REQUEST)
        if archivo.size > 2 * 1024 * 1024:
            return Response({'detail': 'El CSV pasa de 2 MB.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            pagos, omitidos, errores = leer_movimientos(archivo.read())
        except ValueError as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response({'pagos': revisar_pagos(pagos), 'omitidos': omitidos, 'errores': errores})

    @action(detail=False, methods=['post'])
    def guardar(self, request):
        """Guarda la lista de una vez. Los repetidos (misma clave que un pago
        activo o que otro de la lista) no entran y se devuelven aparte."""
        lista = request.data.get('pagos')
        if not isinstance(lista, list) or not lista:
            # Vacía es un fallo del cliente, no "0 guardados": que se vea.
            return Response({'detail': 'No llegó ningún pago que guardar.'},
                            status=status.HTTP_400_BAD_REQUEST)
        datos = PagoSerializer(data=lista, many=True)
        datos.is_valid(raise_exception=True)
        usuario = request.user.username
        guardados, repetidos = [], []
        with transaction.atomic(using=get_db_alias()):
            vistas = set(Pago.objects.filter(estado='activa').exclude(clave_rastreo='')
                         .values_list('clave_rastreo', flat=True))
            for v in datos.validated_data:
                if v.get('clave_rastreo') and v['clave_rastreo'] in vistas:
                    repetidos.append(v['clave_rastreo'])
                    continue
                vistas.add(v.get('clave_rastreo'))
                guardados.append(Pago.objects.create(**v, created_by=usuario, updated_by=usuario))
        return Response({'guardados': len(guardados), 'repetidos': repetidos},
                        status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['get'], url_path='por-aplicar')
    def por_aplicar(self, request):
        """Los pagos activos con algo por aplicar, del más viejo al más
        reciente: la lista numerada de la que se elige el pago de cada abono."""
        qs = (Pago.objects.filter(estado='activa')
              .annotate(aplicado=Coalesce(Sum('abonos__monto', filter=Q(abonos__estado='activa')),
                                          Decimal('0')))
              .filter(aplicado__lt=F('monto')).order_by('fecha', 'id'))
        return Response(PagoSerializer(qs, many=True).data)

    @action(detail=True, methods=['post'])
    def cancelar(self, request, pk=None):
        if not request.user.is_staff:
            raise PermissionDenied('Solo un administrador puede cancelar pagos.')
        usuario = request.user.username
        with transaction.atomic(using=get_db_alias()):
            pago = Pago.objects.select_for_update().get(pk=pk)
            if pago.estado == 'cancelada':
                return Response({'detail': 'El pago ya está cancelado.'},
                                status=status.HTTP_400_BAD_REQUEST)
            for abono in pago.abonos.select_for_update().filter(estado='activa'):
                factura = Factura.objects.select_for_update().get(pk=abono.factura_id)
                abono.estado, abono.updated_by = 'cancelada', usuario
                abono.save(update_fields=['estado', 'updated_by', 'updated_at'])
                mover_abonado(factura, -abono.monto, pago, usuario)
            pago.estado, pago.updated_by = 'cancelada', usuario
            pago.save(update_fields=['estado', 'updated_by', 'updated_at'])
        return Response(self.get_serializer(pago).data)

    @action(detail=False, methods=['get'])
    def bancos(self, request):
        return Response(list(Banco.objects.values_list('nombre', flat=True)))


class AbonoSerializer(serializers.ModelSerializer):
    monto = serializers.DecimalField(max_digits=14, decimal_places=2, min_value=Decimal('0.01'))
    parcialidad = serializers.IntegerField(min_value=1, max_value=999)
    fecha = serializers.DateField(source='pago.fecha', read_only=True)
    clave_rastreo = serializers.CharField(source='pago.clave_rastreo', read_only=True)
    banco_nombre = serializers.CharField(source='pago.banco', read_only=True)

    class Meta:
        model = Abono
        fields = ['id', 'pago', 'factura', 'monto', 'parcialidad', 'estado', 'fecha',
                  'clave_rastreo', 'banco_nombre', 'created_by', 'created_at']
        read_only_fields = ['estado', 'created_by', 'created_at']

    def validate_pago(self, pago):
        if pago.estado != 'activa':
            raise serializers.ValidationError('El pago está cancelado.')
        return pago

    def validate_factura(self, factura):
        # Salen en el buscador pero no se abonan (usuario, 2026-10-22).
        if factura.estado != 'activa':
            raise serializers.ValidationError('La factura está cancelada.')
        if factura.total is None:
            raise serializers.ValidationError('La factura aún no tiene XML: no hay total.')
        return factura


class AbonoViewSet(_SoloFinanzas, mixins.ListModelMixin, mixins.CreateModelMixin,
                   viewsets.GenericViewSet):
    """COMPENSACIÓN, lo que un pago abona a una factura. Cancelar, solo staff."""
    serializer_class = AbonoSerializer

    def get_queryset(self):
        qs = Abono.objects.select_related('pago')
        if self.action == 'list':
            # Siempre los de UNA factura: no hay pantalla con todos.
            qs = qs.filter(factura_id=self.request.query_params.get('factura') or 0)
        return qs

    def _respuesta(self, abono, factura, pago, **kw):
        return Response({'abono': self.get_serializer(abono).data,
                         'factura': _fila_al_dia(factura),
                         'pago': PagoSerializer(pago).data}, **kw)

    def create(self, request, *args, **kwargs):
        datos = self.get_serializer(data=request.data)
        datos.is_valid(raise_exception=True)
        with transaction.atomic(using=get_db_alias()):
            factura = Factura.objects.select_for_update().get(pk=datos.validated_data['factura'].pk)
            abono = datos.save(created_by=request.user.username, updated_by=request.user.username)
            mover_abonado(factura, abono.monto, abono.pago, request.user.username)
        return self._respuesta(abono, factura, abono.pago, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'])
    def cancelar(self, request, pk=None):
        if not request.user.is_staff:
            raise PermissionDenied('Solo un administrador puede cancelar abonos.')
        with transaction.atomic(using=get_db_alias()):
            abono = Abono.objects.select_for_update().select_related('pago').get(pk=pk)
            if abono.estado == 'cancelada':
                return Response({'detail': 'El abono ya está cancelado.'},
                                status=status.HTTP_400_BAD_REQUEST)
            factura = Factura.objects.select_for_update().get(pk=abono.factura_id)
            abono.estado, abono.updated_by = 'cancelada', request.user.username
            abono.save(update_fields=['estado', 'updated_by', 'updated_at'])
            mover_abonado(factura, -abono.monto, abono.pago, request.user.username)
        return self._respuesta(abono, factura, abono.pago)


class FacturaExcelSerializer(serializers.ModelSerializer):
    class Meta:
        model = FacturaExcel
        fields = ['id', 'empresa', 'serie', 'folio', 'nombre', 'rfc',
                  'fecha_emision', 'total', 'estado']
        read_only_fields = fields


class FacturaExcelViewSet(CancelableMixin, viewsets.GenericViewSet):
    """AUDITORÍA: el Excel y su comparación con Facturación. Solo admin
    (usuario, 2026-10-14)."""
    serializer_class = FacturaExcelSerializer
    queryset = FacturaExcel.objects.all()
    pagination_class = PaginaDeCien

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not request.user.is_staff:
            raise PermissionDenied('Auditoría es solo para administradores.')

    @action(detail=False, methods=['get'])
    def comparacion(self, request):
        """La página de 100 de la comparación, con ?estado= y los mismos
        filtros que Facturación. La alerta va aparte y entera: cuenta TODAS las
        que no cuadran, no solo las de la página."""
        datos, p = comparar(), request.query_params
        filas = [f for f in datos['facturas'] if p.get('estado') in (None, '', f['estado'])]
        respuesta = self.get_paginated_response(self.paginate_queryset(filtrar_tabla(filas, p)))
        respuesta.data['distintas'] = [f for f in datos['facturas'] if f['diferencias']]
        respuesta.data['solo_en_facturacion'] = datos['solo_en_facturacion']
        return respuesta

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser])
    def cargar(self, request):
        archivo = request.FILES.get('archivo')
        if not archivo:
            return Response({'detail': 'Falta el archivo.'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            resumen = cargar_facturas(archivo, request.user.username)
        except ValueError as e:
            return Response({'detail': str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(resumen, status=status.HTTP_201_CREATED)
